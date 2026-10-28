import re
import shutil
from pathlib import Path
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE_DIR = Path(__file__).resolve().parent
FIG_DIR = BASE_DIR / 'Gambar_Skripsi'
FIG_DIR.mkdir(exist_ok=True)


def generate_kerangka_pemikiran_png() -> Path:
    """
    Membuat diagram alir Gambar 1 (Bagan Kerangka Pemikiran) beresolusi 300 DPI
    agar di dalam dokumen Word (.docx) tampil sebagai bagan visual yang rapi,
    bukan blok kode teks Mermaid.
    """
    fig, ax = plt.subplots(figsize=(10.5, 8.2), dpi=300)
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 100)
    ax.axis('off')

    def draw_box(x, y, w, h, title, lines, fc='#eef4fa', ec='#1f4e78', fontsize=8.0):
        box = mpatches.FancyBboxPatch(
            (x, y), w, h,
            boxstyle="round,pad=0.6,rounding_size=1.2",
            facecolor=fc, edgecolor=ec, linewidth=1.3
        )
        ax.add_patch(box)
        text_str = f"{title}\n" + "\n".join(lines) if lines else title
        ax.text(
            x + w / 2, y + h / 2, text_str,
            ha='center', va='center', fontsize=fontsize,
            color='#111111', family='DejaVu Sans',
            linespacing=1.32
        )

    def draw_arrow(x1, y1, x2, y2):
        ax.annotate(
            '', xy=(x2, y2), xytext=(x1, y1),
            arrowprops=dict(arrowstyle='->', color='#1f4e78', lw=1.5, shrinkA=0, shrinkB=0)
        )

    draw_box(
        8, 86, 84, 11.5,
        "MASALAH PENELITIAN",
        [
            "1. Diperlukan pemodelan prediktif anomali termal perairan Danau Maninjau berbasis GEE dengan pengendalian lookahead bias.",
            "2. Perlu diaudit representasi fisik grid ERA5-Land (11,1 km) serta kesesuaiannya terhadap satelit MODIS & Landsat 8/9."
        ],
        fc='#fdf2e9', ec='#b95000'
    )
    draw_arrow(50, 86, 50, 81.5)

    draw_box(
        8, 63.5, 84, 17.5,
        "EKSTRAKSI & AUDIT SPASIAL-OBSERVASIONAL GEE (2016–2025, 3.653 HARI)",
        [
            "• Prediktor & Target Reanalisis: ERA5-Land (0,1° ≈ 11,1 km) & CHIRPS (0,05° ≈ 5,5 km) pada ROI Inti Air (r = 2 km, JRC > 90%)",
            "• Audit Spasial Grid Utama ERA5-Land: Fraksi air permanen JRC = 52,20% (64,54 km²) & stratifikasi modul FLake",
            "• Evaluasi Sensitivitas Satelit: MODIS Terra/Aqua (≥15/22 piksel, std ≤ 0,5°C, |angle| ≤ 30°) & Landsat 8/9",
            "• Penetapan era5_lake_mix_temp_c sebagai target reanalisis harian dengan batasan validitas observasional eksplisit"
        ],
        fc='#eaf2f8', ec='#1f618d'
    )
    draw_arrow(50, 63.5, 50, 59.0)

    draw_box(
        8, 42.5, 84, 16.0,
        "PREPROCESSING PENGENDALIAN KEBOCORAN INFORMASI & ROLLING-ORIGIN CV",
        [
            "1. Indeks Kalender Konsisten 365 Hari (cal_doy_365) tanpa pergeseran pasca-Februari pada tahun kabisat",
            "2. Isolasi perhitungan klimatologi hanya dari data latih masing-masing fold & pemisahan kronologis tanggal target (t+h ≤ t_akhir)",
            "3. Dua Rezim Fitur Input (≤ t): Skema Utama (73 Fitur: 18 Termal + 55 Meteo) & Skema Ablasi Murni Meteo (55 Fitur)",
            "4. Seleksi Hyperparameter via 3-Fold Rolling-Origin CV (2016–2023) & Uji Independen Satu Kali (2024–2025)"
        ],
        fc='#e8f8f5', ec='#117864'
    )

    model_boxes = [
        (4, 27.5, 21.5, 10.5, "Dua Baseline\nPersistence", ["Anomali: Â_t+h = A_t", "Suhu: T̂_t+h = T_t"], '#f2f4f4', '#566573'),
        (28, 27.5, 21.0, 10.5, "Random Forest\nRegressor (RF)", ["Ensemble Bagging", "Pohon Keputusan"], '#ebf5fb', '#2874a6'),
        (51, 27.5, 21.0, 10.5, "XGBoost\nRegressor", ["Gradient Boosting", "Sekuensial Terregularisasi"], '#ebf5fb', '#2874a6'),
        (74, 27.5, 22.0, 10.5, "Support Vector\nRegression (SVR)", ["Kernel RBF +", "StandardScaler Pipeline"], '#ebf5fb', '#2874a6'),
    ]
    for bx, by, bw, bh, btitle, blines, bfc, bec in model_boxes:
        draw_box(bx, by, bw, bh, btitle, blines, fc=bfc, ec=bec, fontsize=7.6)
        draw_arrow(bx + bw / 2, 42.5, bx + bw / 2, by + bh)
        draw_arrow(bx + bw / 2, by, bx + bw / 2, 23.0)

    draw_box(
        8, 10.5, 84, 12.0,
        "EVALUASI KINERJA PREDIKSI & ANALISIS DESKRIPTIF CATATAN UPWELLING",
        [
            "1. Komparasi MAE, RMSE (Metrik Utama), dan R² pada Data Uji Independen (2024–2025) untuk Horizon t+1 dan t+3",
            "2. Analisis Kontribusi Riwayat Suhu vs Kemampuan Murni Meteorologi (Feature Importance & Uji Ablasi 73 vs 55 Fitur)",
            "3. Evaluasi Deskriptif Anomali Termal Reanalisis pada 15 Kandidat Catatan Historis (Exact_Day, Multi_Day, Month_Only)"
        ],
        fc='#fef9e7', ec='#b7950b'
    )
    draw_arrow(50, 10.5, 50, 7.0)

    draw_box(
        12, 0.8, 76, 6.0,
        "HASIL AKHIR (OUTPUT RISET SKEMA 1)",
        ["Komparasi Kinerja Algoritma pada Target Reanalisis ERA5-Land & Dokumentasi Ilmiah Keterbatasan Validasi Satelit/Upwelling"],
        fc='#f4ecf7', ec='#6c3483', fontsize=8.0
    )

    plt.tight_layout()
    out_sub = FIG_DIR / 'Gambar_0_Kerangka_Pemikiran.png'
    out_root = BASE_DIR / 'Gambar_0_Kerangka_Pemikiran.png'
    fig.savefig(out_sub, dpi=300, bbox_inches='tight')
    fig.savefig(out_root, dpi=300, bbox_inches='tight')
    plt.close()
    return out_sub


def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key in ["val", "color", "sz", "space"]:
                if key in edge_data:
                    element.set(qn('w:{}'.format(key)), str(edge_data[key]))


def set_cell_shading(cell, fill_hex: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)


def clean_latex_and_md(text: str) -> str:
    s = text.replace('<br><br>', '\n').replace('<br/>', '\n').replace('<br>', '\n')
    replacements = [
        (r'\blacksquare', '■'),
        (r'\checkmark', '✓'),
        (r'\le', '≤'),
        (r'\ge', '≥'),
        (r'\approx', '≈'),
        (r'\equiv', '≡'),
        (r'\pm', '±'),
        (r'\times', '×'),
        (r'\cdot', '·'),
        (r'\rightarrow', '→'),
        (r'\dots', '...'),
        (r'\Delta', 'Δ'),
        (r'\varepsilon', 'ε'),
        (r'\eta', 'η'),
        (r'\gamma', 'γ'),
        (r'\lambda', 'λ'),
        (r'\rho', 'ρ'),
        (r'\phi', 'φ'),
        (r'\mu', 'μ'),
        (r'\circ', '°'),
        (r'^\circ\text{C}', '°C'),
        (r'^\circ', '°'),
        (r'\text{H}_2\text{S}', 'H₂S'),
        (r'\text{NH}_3', 'NH₃'),
        (r'\text{km}^2', 'km²'),
        (r'\text{m}^2', 'm²'),
        (r'R^2', 'R²'),
        (r'H_1', 'H₁'),
        (r'H_2', 'H₂'),
        (r'H_3', 'H₃'),
        (r't_0', 't₀'),
        (r'R_1', 'R₁'),
        (r'R_2', 'R₂'),
        (r'L_1', 'L₁'),
        (r'L_2', 'L₂'),
        (r'u_{10}', 'u₁₀'),
        (r'v_{10}', 'v₁₀'),
        (r'\hat{A}_{t+h}^{\text{pers-A}}', 'Â_(t+h)^(pers-A)'),
        (r'\hat{T}_{t+h}^{\text{pers-A}}', 'T̂_(t+h)^(pers-A)'),
        (r'\hat{T}_{t+h}^{\text{pers-T}}', 'T̂_(t+h)^(pers-T)'),
        (r'\hat{A}_{t+h}^{\text{pers-T}}', 'Â_(t+h)^(pers-T)'),
        (r'\bar{T}_{\text{clim}, d(t+h)}', 'T̄_clim,d(t+h)'),
        (r'\bar{T}_{\text{clim}, d(t)}', 'T̄_clim,d(t)'),
        (r'\bar{T}_{\text{clim}, d}', 'T̄_clim,d'),
        (r'\hat{A}_{t+h}', 'Â_(t+h)'),
        (r'\hat{T}_{t+h}', 'T̂_(t+h)'),
        (r'\hat{T}_{t+1}', 'T̂_(t+1)'),
        (r'\hat{A}_{t+1}', 'Â_(t+1)'),
        (r'\hat{T}_{t+3}', 'T̂_(t+3)'),
        (r'\hat{A}_{t+3}', 'Â_(t+3)'),
        (r'A_{t+h}', 'A_(t+h)'),
        (r'T_{t+h}', 'T_(t+h)'),
        (r'y_{t+1}', 'y_(t+1)'),
        (r'y_{t+3}', 'y_(t+3)'),
        (r'A_t', 'A_t'),
        (r'T_t', 'T_t'),
        (r'\hat{y}', 'ŷ'),
        (r'\bar{y}', 'ȳ'),
    ]
    for old, new in replacements:
        s = s.replace(old, new)

    s = re.sub(r'\\text\{([^}]*)\}', r'\1', s)
    s = re.sub(r'\\mathbf\{([^}]*)\}', r'\1', s)
    s = s.replace(r'\left', '').replace(r'\right', '').replace(r'\,', ' ')
    s = re.sub(r'\$([^$]+)\$', r'\1', s)
    return s


def add_formatted_runs(paragraph, text: str, base_font_size=12, base_italic=False, base_bold=False):
    cleaned = clean_latex_and_md(text)
    pattern = re.compile(r'(\*\*\*[^*]+\*\*\*|\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)')
    tokens = pattern.split(cleaned)
    for tok in tokens:
        if not tok:
            continue
        run = paragraph.add_run()
        run.font.name = 'Times New Roman'
        run.font.size = Pt(base_font_size)
        run.font.color.rgb = RGBColor(0, 0, 0)
        if tok.startswith('***') and tok.endswith('***'):
            run.text = tok[3:-3]
            run.bold = True
            run.italic = True
        elif tok.startswith('**') and tok.endswith('**'):
            run.text = tok[2:-2]
            run.bold = True
            run.italic = base_italic
        elif tok.startswith('*') and tok.endswith('*'):
            run.text = tok[1:-1]
            run.bold = base_bold
            run.italic = True
        elif tok.startswith('`') and tok.endswith('`'):
            run.text = tok[1:-1]
            run.font.name = 'Consolas'
            run.font.size = Pt(max(base_font_size - 1.5, 8.0))
            run.bold = base_bold
            run.italic = base_italic
        else:
            run.text = tok
            run.bold = base_bold
            run.italic = base_italic


def add_front_matter(doc, is_sempro: bool):
    """
    Menambahkan Halaman Sampul, Lembar Persetujuan/Pengesahan, dan Daftar Isi
    sesuai panduan skripsi FILKOM Universitas Djuanda.
    """
    doc_label = "PROPOSAL SKRIPSI" if is_sempro else "SKRIPSI"
    stage_label = "Seminar Proposal Skripsi (Sempro)" if is_sempro else "Seminar Hasil Skripsi (Semhas) / Sidang Akhir"

    # 1. Halaman Sampul
    p_top = doc.add_paragraph()
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_top.paragraph_format.space_before = Pt(12)
    p_top.paragraph_format.space_after = Pt(18)
    p_top.paragraph_format.line_spacing = 1.5
    add_formatted_runs(p_top, f"**{doc_label}**", base_font_size=14, base_bold=True)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(24)
    p_title.paragraph_format.line_spacing = 1.5
    add_formatted_runs(
        p_title,
        "**KOMPARASI ALGORITMA RANDOM FOREST, XGBOOST, DAN SUPPORT VECTOR REGRESSION "
        "UNTUK PREDIKSI ANOMALI SUHU LAPISAN TERCAMPUR REANALISIS ERA5-LAND "
        "DI KAWASAN DANAU MANINJAU BERBASIS GOOGLE EARTH ENGINE**",
        base_font_size=13, base_bold=True
    )

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(12)
    p_sub.paragraph_format.space_after = Pt(36)
    p_sub.paragraph_format.line_spacing = 1.5
    add_formatted_runs(
        p_sub,
        f"Diajukan untuk Memenuhi Salah Satu Syarat Mengikuti {stage_label}\n"
        "pada Program Studi Ilmu Komputer Jenjang Strata Satu (S1)\n"
        "Skema 1: Pengembangan Algoritma / Model / Analisis",
        base_font_size=12
    )

    p_auth = doc.add_paragraph()
    p_auth.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_auth.paragraph_format.space_before = Pt(24)
    p_auth.paragraph_format.space_after = Pt(48)
    p_auth.paragraph_format.line_spacing = 1.5
    add_formatted_runs(
        p_auth,
        "Disusun Oleh:\n**[NAMA MAHASISWA]**\n**NIM: [NIM MAHASISWA]**",
        base_font_size=12
    )

    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_before = Pt(36)
    p_inst.paragraph_format.space_after = Pt(0)
    p_inst.paragraph_format.line_spacing = 1.5
    add_formatted_runs(
        p_inst,
        "**PROGRAM STUDI ILMU KOMPUTER**\n"
        "**FAKULTAS ILMU KOMPUTER**\n"
        "**UNIVERSITAS DJUANDA**\n"
        "**BOGOR**\n"
        "**2026**",
        base_font_size=13, base_bold=True
    )

    # 2. Lembar Persetujuan Pembimbing
    doc.add_page_break()
    p_ap_title = doc.add_paragraph()
    p_ap_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ap_title.paragraph_format.space_after = Pt(18)
    add_formatted_runs(
        p_ap_title,
        f"**LEMBAR PERSETUJUAN {doc_label}**",
        base_font_size=14, base_bold=True
    )

    p_ap_body = doc.add_paragraph()
    p_ap_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_ap_body.paragraph_format.line_spacing = 1.5
    p_ap_body.paragraph_format.space_after = Pt(24)
    add_formatted_runs(
        p_ap_body,
        f"Naskah {doc_label.title()} dengan judul **\"Komparasi Algoritma Random Forest, XGBoost, dan Support Vector Regression "
        f"untuk Prediksi Anomali Suhu Lapisan Tercampur Reanalisis ERA5-Land di Kawasan Danau Maninjau Berbasis Google Earth Engine\"** "
        f"yang disusun oleh **[Nama Mahasiswa]** (NIM: **[NIM Mahasiswa]**) telah diperiksa dan disetujui untuk diajukan pada "
        f"**{stage_label}** Program Studi Ilmu Komputer, Fakultas Ilmu Komputer, Universitas Djuanda.",
        base_font_size=12
    )

    tbl_sig = doc.add_table(rows=2, cols=2)
    tbl_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_labels = [
        ("Dosen Pembimbing Utama,\n\n\n\n**[Nama Dosen Pembimbing I, Gelar]**\nNIDN: [NIDN Pembimbing I]",
         "Dosen Pembimbing Pendamping,\n\n\n\n**[Nama Dosen Pembimbing II, Gelar]**\nNIDN: [NIDN Pembimbing II]"),
        ("\nMengetahui,\nKetua Program Studi Ilmu Komputer\n\n\n\n**[Nama Kaprodi Ilmu Komputer, Gelar]**\nNIDN: [NIDN Kaprodi]",
         "\nMenyetujui,\nDekan Fakultas Ilmu Komputer\n\n\n\n**[Nama Dekan FILKOM, Gelar]**\nNIDN: [NIDN Dekan]")
    ]
    for r_i in range(2):
        for c_i in range(2):
            c_cell = tbl_sig.rows[r_i].cells[c_i]
            p_c = c_cell.paragraphs[0]
            p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_c.paragraph_format.line_spacing = 1.15
            add_formatted_runs(p_c, sig_labels[r_i][c_i], base_font_size=11)

    # 3. Halaman Daftar Isi Kerangka
    doc.add_page_break()
    p_toc_h = doc.add_paragraph()
    p_toc_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_toc_h.paragraph_format.space_after = Pt(14)
    add_formatted_runs(p_toc_h, "**DAFTAR ISI**", base_font_size=14, base_bold=True)

    toc_items = [
        "HALAMAN SAMPUL",
        f"LEMBAR PERSETUJUAN {doc_label}",
        "DAFTAR ISI",
    ]
    if not is_sempro:
        toc_items.extend([
            "ABSTRAK",
            "DAFTAR TABEL DAN DAFTAR GAMBAR",
        ])
    else:
        toc_items.append("DAFTAR TABEL DAN DAFTAR GAMBAR PROPOSAL")

    toc_items.extend([
        "BAB I — PENDAHULUAN (Latar Belakang, Permasalahan, Tujuan, Spesifikasi Produk Skema 1, Manfaat, Batasan, Definisi Operasional, Sistematika)",
        "BAB II — TINJAUAN PUSTAKA DAN KERANGKA TEORITIS (Objek Penelitian Danau Maninjau, 10 Penelitian Terdahulu 2022–2026, Landasan Teori, Kerangka Pemikiran, Hipotesis H₁–H₃)",
        "BAB III — METODOLOGI PENELITIAN (Tahapan Skema 1, Arsitektur Model, Prosedur GEE & Audit Spasial-Satelit, Preprocessing Pengendalian Kebocoran Informasi 73 vs 55 Fitur, Desain 3-Fold ROCV & Evaluasi, Jadwal Penelitian)",
    ])
    if not is_sempro:
        toc_items.extend([
            "BAB IV — HASIL DAN PEMBAHASAN (Implementasi Pipa GEE–Python, Audit Grid ERA5-Land, Sensitivitas ERA5 vs MODIS/Landsat, Hasil Tuning ROCV, Komparasi Uji Independen 2024–2025, Feature Importance, Evaluasi Eksploratif Catatan Upwelling, Pembahasan)",
            "BAB V — KESIMPULAN DAN SARAN (Kesimpulan Rumusan Masalah 1–2 & Hipotesis H₁–H₃, Saran Pengembangan)",
        ])
    toc_items.extend([
        "DAFTAR RUJUKAN (11 Artikel Jurnal Primer & 7 Berita Resmi LKBN ANTARA)",
        "LAMPIRAN PENDUKUNG (Daftar Artefak Kode/Dataset & Dokumentasi Metadata Bibliografi)",
    ])
    for item in toc_items:
        p_t = doc.add_paragraph()
        p_t.paragraph_format.line_spacing = 1.35
        p_t.paragraph_format.space_before = Pt(1)
        p_t.paragraph_format.space_after = Pt(2)
        add_formatted_runs(p_t, f"•  {item}", base_font_size=11.5)

    if is_sempro:
        doc.add_page_break()
        p_lt = doc.add_paragraph()
        p_lt.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_lt.paragraph_format.space_after = Pt(10)
        add_formatted_runs(p_lt, "**DAFTAR TABEL DAN DAFTAR GAMBAR PROPOSAL**", base_font_size=14, base_bold=True)
        prop_lists = [
            "**Tabel 1.** Ringkasan Penelitian Terdahulu (2022–2026)",
            "**Tabel 2.** Tahapan Metodologi Penelitian Pengembangan Model (Skema 1)",
            "**Tabel 3.** Rincian Eksak Konstruksi Fitur Input pada Skema Utama (73 Fitur) dan Skema Ablasi Murni Meteorologi (55 Fitur)",
            "**Tabel 4.** Jadwal Pelaksanaan Penelitian Skripsi",
            "**Gambar 1.** Bagan kerangka pemikiran penelitian komparasi algoritma regresi dan audit spasial-observasional di Danau Maninjau",
        ]
        for pl in prop_lists:
            p_pl = doc.add_paragraph()
            p_pl.paragraph_format.line_spacing = 1.35
            add_formatted_runs(p_pl, f"•  {pl}", base_font_size=11.5)


def render_markdown_to_docx(md_lines: list, out_docx: Path, kerangka_png: Path, is_sempro: bool):
    doc = Document()

    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(4.0)
        section.left_margin = Cm(4.0)
        section.bottom_margin = Cm(3.0)
        section.right_margin = Cm(3.0)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)

    add_front_matter(doc, is_sempro=is_sempro)

    in_code_block = False
    code_lang = ''
    code_lines = []
    table_lines = []
    in_abstract = False
    in_references = False

    def flush_table(tbl_lines):
        if len(tbl_lines) < 2:
            return
        rows_data = []
        for tl in tbl_lines:
            parts = [c.strip() for c in tl.strip().strip('|').split('|')]
            if all(re.match(r'^:?-{2,}:?$', p) for p in parts if p):
                continue
            rows_data.append(parts)
        if not rows_data:
            return
        num_cols = max(len(r) for r in rows_data)
        tbl = doc.add_table(rows=len(rows_data), cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = True

        # Gunakan ukuran font lebih ringkas (8.5 pt) untuk tabel lebar (>= 7 kolom)
        tbl_font_sz = 8.5 if num_cols >= 7 else 9.5

        border_spec = {"val": "single", "sz": "4", "space": "0", "color": "333333"}
        for r_idx, r_vals in enumerate(rows_data):
            row_cells = tbl.rows[r_idx].cells
            for c_idx in range(num_cols):
                val = r_vals[c_idx] if c_idx < len(r_vals) else ''
                cell = row_cells[c_idx]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_border(
                    cell,
                    top=border_spec, bottom=border_spec,
                    left=border_spec, right=border_spec
                )
                if r_idx == 0:
                    set_cell_shading(cell, "E8EEF5")
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(1.5)
                p.paragraph_format.space_after = Pt(1.5)
                p.paragraph_format.line_spacing = 1.0
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if r_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
                add_formatted_runs(p, val, base_font_size=tbl_font_sz, base_bold=(r_idx == 0))

        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_before = Pt(0)
        spacer.paragraph_format.space_after = Pt(4)

    i = 0
    while i < len(md_lines):
        raw_line = md_lines[i]
        stripped = raw_line.strip()

        if stripped.startswith('```'):
            if table_lines:
                flush_table(table_lines)
                table_lines = []
            if not in_code_block:
                in_code_block = True
                code_lang = stripped[3:].strip().lower()
                code_lines = []
            else:
                in_code_block = False
                if code_lang == 'mermaid':
                    p_img = doc.add_paragraph()
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img.paragraph_format.space_before = Pt(6)
                    p_img.paragraph_format.space_after = Pt(4)
                    run_img = p_img.add_run()
                    run_img.add_picture(str(kerangka_png), width=Cm(14.0))
                else:
                    p_code = doc.add_paragraph()
                    p_code.paragraph_format.left_indent = Cm(0.5)
                    p_code.paragraph_format.right_indent = Cm(0.5)
                    p_code.paragraph_format.space_before = Pt(4)
                    p_code.paragraph_format.space_after = Pt(6)
                    p_code.paragraph_format.line_spacing = 1.0
                    run_c = p_code.add_run("\n".join(code_lines))
                    run_c.font.name = 'Consolas'
                    run_c.font.size = Pt(9.0)
                    run_c.font.color.rgb = RGBColor(20, 20, 20)
                code_lang = ''
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(raw_line)
            i += 1
            continue

        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines.append(stripped)
            i += 1
            continue
        elif table_lines:
            flush_table(table_lines)
            table_lines = []

        if stripped in ('---', '') or stripped in (
            '# NASKAH LENGKAP SKRIPSI (BAB I – BAB V)',
            '# PROPOSAL SKRIPSI',
        ):
            i += 1
            continue

        # Skip duplicate front-matter metadata lines already placed on the cover page
        if stripped.startswith('**JUDUL SKRIPSI:**') or stripped.startswith('**Skema Penelitian:**') or stripped.startswith('**Program Studi:**'):
            i += 1
            continue

        if stripped.startswith('$$') and stripped.endswith('$$'):
            eq_text = stripped[2:-2].strip()
            eq_text = (
                eq_text.replace(r'\text{SSE} = \sum_{i \in R_1} (y_i - \bar{y}_{R_1})^2 + \sum_{i \in R_2} (y_i - \bar{y}_{R_2})^2',
                                'SSE = ∑_(i ∈ R₁) (y_i − ȳ_R₁)² + ∑_(i ∈ R₂) (y_i − ȳ_R₂)²')
                .replace(r'\hat{y} = \frac{1}{K} \sum_{k=1}^{K} h_k(\mathbf{x})',
                         'ŷ = (1 / K) ∑_(k=1..K) h_k(x)')
                .replace(r'\hat{y}_i^{(m)} = \hat{y}_i^{(m-1)} + \eta f_m(\mathbf{x}_i)',
                         'ŷ_i^(m) = ŷ_i^(m-1) + η f_m(x_i)')
                .replace(r'\mathcal{L}^{(m)} = \sum_{i=1}^{n} \frac{1}{2}\left(y_i - \hat{y}_i^{(m)}\right)^2 + \gamma T + \frac{1}{2}\lambda \sum_{j=1}^{T} w_j^2',
                         'L^(m) = ∑_(i=1..n) ½ (y_i − ŷ_i^(m))² + γT + ½ λ ∑_(j=1..T) w_j²')
                .replace(r'\min_{\mathbf{w}, b, \xi, \xi^*} \frac{1}{2} \|\mathbf{w}\|^2 + C \sum_{i=1}^{n} (\xi_i + \xi_i^*)',
                         'min_(w,b,ξ,ξ*) ½ ||w||² + C ∑_(i=1..n) (ξ_i + ξ_i*)')
                .replace(r'A_t = T_t - \bar{T}_{\text{clim}, d(t)}',
                         'A_t = T_t − T̄_clim,d(t)')
            )
            p_eq = doc.add_paragraph()
            p_eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_eq.paragraph_format.space_before = Pt(4)
            p_eq.paragraph_format.space_after = Pt(4)
            p_eq.paragraph_format.line_spacing = 1.15
            add_formatted_runs(p_eq, eq_text, base_font_size=11.5, base_italic=True)
            i += 1
            continue

        img_match = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)$', stripped)
        if img_match:
            img_rel = img_match.group(2)
            img_path = BASE_DIR / img_rel
            if not img_path.is_file():
                img_path = BASE_DIR / Path(img_rel).name
            if img_path.is_file():
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(8)
                p_img.paragraph_format.space_after = Pt(4)
                run_img = p_img.add_run()
                run_img.add_picture(str(img_path), width=Cm(14.0))
            i += 1
            continue

        if stripped.startswith('# '):
            heading_text = stripped[2:].strip()
            in_abstract = False
            in_references = ('DAFTAR RUJUKAN' in heading_text)
            doc.add_page_break()
            p_h1 = doc.add_paragraph()
            p_h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_h1.paragraph_format.space_before = Pt(0)
            p_h1.paragraph_format.space_after = Pt(12)
            p_h1.paragraph_format.line_spacing = 1.5
            if ' — ' in heading_text:
                parts_h = heading_text.split(' — ', 1)
                add_formatted_runs(p_h1, parts_h[0] + "\n" + parts_h[1], base_font_size=14, base_bold=True)
            else:
                add_formatted_runs(p_h1, heading_text, base_font_size=14, base_bold=True)
            i += 1
            continue

        if stripped.startswith('## '):
            sec_text = stripped[3:].strip()
            if sec_text == 'ABSTRAK':
                in_abstract = True
                in_references = False
                doc.add_page_break()
                p_sec = doc.add_paragraph()
                p_sec.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_sec.paragraph_format.space_before = Pt(0)
                p_sec.paragraph_format.space_after = Pt(10)
                add_formatted_runs(p_sec, sec_text, base_font_size=14, base_bold=True)
            elif sec_text == 'DAFTAR TABEL DAN DAFTAR GAMBAR':
                in_abstract = False
                doc.add_page_break()
                p_sec = doc.add_paragraph()
                p_sec.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_sec.paragraph_format.space_before = Pt(0)
                p_sec.paragraph_format.space_after = Pt(10)
                add_formatted_runs(p_sec, sec_text, base_font_size=14, base_bold=True)
            else:
                in_abstract = False
                p_sec = doc.add_paragraph()
                p_sec.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p_sec.paragraph_format.space_before = Pt(10)
                p_sec.paragraph_format.space_after = Pt(4)
                p_sec.paragraph_format.line_spacing = 1.5
                add_formatted_runs(p_sec, sec_text, base_font_size=12, base_bold=True)
            i += 1
            continue

        if stripped.startswith('### ') or stripped.startswith('#### '):
            lvl = 4 if stripped.startswith('#### ') else 3
            sub_text = stripped[lvl + 1:].strip()
            p_sub = doc.add_paragraph()
            p_sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_sub.paragraph_format.space_before = Pt(8)
            p_sub.paragraph_format.space_after = Pt(3)
            p_sub.paragraph_format.line_spacing = 1.5
            add_formatted_runs(p_sub, sub_text, base_font_size=12, base_bold=True)
            i += 1
            continue

        if (stripped.startswith('**Tabel ') or stripped.startswith('**Gambar ') or
                stripped.startswith('*Sumber:')):
            p_cap = doc.add_paragraph()
            if stripped.startswith('**Gambar ') or stripped.startswith('*Sumber:'):
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_cap.paragraph_format.space_before = Pt(4)
            p_cap.paragraph_format.space_after = Pt(4)
            p_cap.paragraph_format.line_spacing = 1.0
            add_formatted_runs(p_cap, stripped, base_font_size=10.5)
            i += 1
            continue

        bullet_match = re.match(r'^(\*|\d+\.)\s+(.*)$', stripped)
        if bullet_match:
            prefix = bullet_match.group(1)
            item_body = bullet_match.group(2)
            p_li = doc.add_paragraph()
            p_li.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_li.paragraph_format.left_indent = Cm(0.75)
            p_li.paragraph_format.first_line_indent = Cm(-0.45)
            p_li.paragraph_format.space_before = Pt(1.5)
            p_li.paragraph_format.space_after = Pt(1.5)
            p_li.paragraph_format.line_spacing = 1.5
            bullet_label = "•  " if prefix == '*' else f"{prefix} "
            add_formatted_runs(p_li, bullet_label + item_body, base_font_size=12)
            i += 1
            continue

        sub_bullet_match = re.match(r'^\s+(\*|\-)\s+(.*)$', raw_line)
        if sub_bullet_match:
            item_body = sub_bullet_match.group(2)
            p_sli = doc.add_paragraph()
            p_sli.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_sli.paragraph_format.left_indent = Cm(1.35)
            p_sli.paragraph_format.first_line_indent = Cm(-0.45)
            p_sli.paragraph_format.space_before = Pt(1)
            p_sli.paragraph_format.space_after = Pt(1)
            p_sli.paragraph_format.line_spacing = 1.5
            add_formatted_runs(p_sli, "–  " + item_body, base_font_size=12)
            i += 1
            continue

        p = doc.add_paragraph()
        if in_abstract:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(6)
        elif in_references and not stripped.startswith('**'):
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Cm(1.0)
            p.paragraph_format.first_line_indent = Cm(-1.0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(6)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(6)
            if not stripped.startswith('**') and not stripped.startswith('*('):
                p.paragraph_format.first_line_indent = Cm(1.0)

        add_formatted_runs(p, stripped, base_font_size=12)
        i += 1

    if table_lines:
        flush_table(table_lines)

    doc.save(out_docx)
    print(f"[SUKSES WORD] Berhasil membuat dokumen: {out_docx}")


def build_all_word_documents():
    kerangka_png = generate_kerangka_pemikiran_png()

    # 1. Dokumen Khusus Seminar Proposal (Sempro): BAB I–III + Daftar Rujukan + Lampiran
    prop_md = BASE_DIR / 'PROPOSAL_SKRIPSI_MANINJAU.md'
    full_md = BASE_DIR / 'NASKAH_LENGKAP_SKRIPSI_MANINJAU_BAB_I_SAMPAI_V.md'
    full_text = full_md.read_text(encoding='utf-8')
    lampiran_idx = full_text.find('# LAMPIRAN PENDUKUNG')
    lampiran_part = "\n\n" + full_text[lampiran_idx:] if lampiran_idx != -1 else ""

    prop_lines = (prop_md.read_text(encoding='utf-8') + lampiran_part).splitlines()
    sempro_docx = BASE_DIR / 'PROPOSAL_SEMPRO_MANINJAU.docx'
    render_markdown_to_docx(prop_lines, sempro_docx, kerangka_png, is_sempro=True)

    # 2. Dokumen Khusus Seminar Hasil (Semhas) / Skripsi Lengkap: BAB I–V + Daftar Rujukan + Lampiran
    full_lines = full_text.splitlines()
    semhas_docx = BASE_DIR / 'NASKAH_SEMHAS_MANINJAU.docx'
    render_markdown_to_docx(full_lines, semhas_docx, kerangka_png, is_sempro=False)

    # Salin juga ke NASKAH_SKRIPSI_MANINJAU_FINAL.docx jika sedang tidak dibuka di Word
    try:
        shutil.copyfile(semhas_docx, BASE_DIR / 'NASKAH_SKRIPSI_MANINJAU_FINAL.docx')
    except PermissionError:
        print("[INFO] NASKAH_SKRIPSI_MANINJAU_FINAL.docx sedang dibuka di Word; gunakan NASKAH_SEMHAS_MANINJAU.docx.")


if __name__ == '__main__':
    build_all_word_documents()
